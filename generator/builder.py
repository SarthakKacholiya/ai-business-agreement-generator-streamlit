import os
from typing import Dict
from docx import Document          # pip install python-docx
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer  # pip install reportlab
from reportlab.lib.styles import getSampleStyleSheet

OUTPUT_DIR = "output"

AVAILABLE_TEMPLATES = ["Service", "Partnership", "NDA"]


def ensure_output_dir() -> None:
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def load_template(name: str) -> str:
    """Loads a template file from /templates."""
    fname = {
        "Service": "service.txt",
        "Partnership": "partnership.txt",
        "NDA": "nda.txt",
    }.get(name, "service.txt")
    path = os.path.join("templates", fname)
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def _clean_lines(text: str) -> str:
    """Collapse extra blank lines produced by missing fields."""
    lines = [line.rstrip() for line in text.splitlines()]
    out = []
    prev_blank = False
    for ln in lines:
        is_blank = (ln.strip() == "")
        if is_blank and prev_blank:
            continue
        out.append(ln)
        prev_blank = is_blank
    return "\n".join(out).strip() + "\n"


def render_template(template: str, data: Dict[str, str]) -> str:
    """
    Safe renderer: replaces {placeholders} with data if present, otherwise blanks.
    """
    class SafeDict(dict):
        def __missing__(self, key):
            return ""
    filled = template.format_map(SafeDict({k: ("" if v is None else str(v)) for k, v in data.items()}))
    return _clean_lines(filled)


def build_txt(text: str, filepath: str) -> str:
    ensure_output_dir()
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(text)
    return filepath


def build_docx(text: str, filepath: str) -> str:
    ensure_output_dir()
    doc = Document()
    for para in text.split("\n"):
        doc.add_paragraph(para)
    doc.save(filepath)
    return filepath


def build_pdf(text: str, filepath: str) -> str:
    ensure_output_dir()
    styles = getSampleStyleSheet()
    story = []
    for para in text.split("\n"):
        story.append(Paragraph(para, styles["Normal"]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(filepath).build(story)
    return filepath
